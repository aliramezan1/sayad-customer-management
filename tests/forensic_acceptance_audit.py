# -*- coding: utf-8 -*-
"""
Forensic Acceptance Audit for Milestone 5:
Independently inspects both generated Excel workbooks and asserts all 4 Acceptance Criteria.
"""

import os
import sys
import openpyxl
from openpyxl.utils import get_column_letter

if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if sys.stderr and hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

ROOT_FILE = r"c:\Users\HP\Desktop\نام و نام خانوادگی مشتریان\گزارش_جامع_اعتباری_مشتریان_صیادی_۱۴۰۵۰۶۱۵.xlsx"
DESKTOP_FILE = r"C:\Users\HP\Desktop\گزارش_جامع_اعتباری_مشتریان_صیادی_۱۴۰۵۰۶۱۵.xlsx"

def audit_file(filepath: str):
    print(f"\n=======================================================")
    print(f"AUDITING FILE: {filepath}")
    print(f"File size: {os.path.getsize(filepath):,} bytes")
    print(f"=======================================================")
    
    wb = openpyxl.load_workbook(filepath, data_only=False)
    wb_data = openpyxl.load_workbook(filepath, data_only=True)
    
    # 1. Structure Audit
    sheets = wb.sheetnames
    print(f"\n[AC 4 - Sheets] Total sheets: {len(sheets)}")
    expected_sheets = [
        "01_خلاصه_مدیریتی",
        "02_مشتریان_یکتا",
        "03_رخدادهای_امروز",
        "04_اقدام_فوری",
        "05_مراقبت",
        "06_بهبود",
        "07_چکهای_تفصیلی",
        "08_مشکلات_هویتی",
        "09_ممیزی",
        "10_راهنما",
    ]
    for s in expected_sheets:
        assert s in sheets, f"Missing sheet: {s}"
        rtl = wb[s].views.sheetView[0].rightToLeft
        assert rtl is True, f"RTL not True for sheet {s}"
        print(f"  - {s}: RTL={rtl}")
    print("  => AC 4 Structure: PASS")
    
    # 2. Sheet 01 Mandatory Statement & KPIs
    ws01 = wb["01_خلاصه_مدیریتی"]
    ws01_data = wb_data["01_خلاصه_مدیریتی"]
    statement_found = False
    statement_text = ""
    for r in range(1, 35):
        for c in range(1, 15):
            val = str(ws01.cell(row=r, column=c).value or "")
            if "محاسبات بانکی بر اساس مشتری یکتا انجام شد" in val:
                statement_found = True
                statement_text = val
                break
        if statement_found:
            break
    print(f"\n[AC 4 - Mandatory Statement]: Found={statement_found}")
    if statement_found:
        print(f"  Statement text: '{statement_text}'")
    assert statement_found, "Mandatory statement NOT FOUND in Sheet 01!"
    
    # 3. Sheet 02 Customer Deduplication Integrity
    ws02 = wb["02_مشتریان_یکتا"]
    ws02_data = wb_data["02_مشتریان_یکتا"]
    
    customers = []
    for r in range(4, 52):
        row_num = ws02_data.cell(row=r, column=1).value
        cid = ws02_data.cell(row=r, column=2).value
        nid = ws02_data.cell(row=r, column=3).value
        name = ws02_data.cell(row=r, column=4).value
        id_status = ws02_data.cell(row=r, column=5).value
        cbi_color = ws02_data.cell(row=r, column=6).value
        chq_cnt = ws02_data.cell(row=r, column=7).value or 0
        fund_amt = ws02_data.cell(row=r, column=8).value or 0
        fund_share = ws02_data.cell(row=r, column=9).value or 0
        in_transit = ws02_data.cell(row=r, column=10).value or 0
        bounced = ws02_data.cell(row=r, column=11).value or 0
        cleared_amt = ws02_data.cell(row=r, column=12).value or 0
        pers_cat = ws02_data.cell(row=r, column=13).value
        pers_periods = ws02_data.cell(row=r, column=14).value or 0
        raw_score = ws02_data.cell(row=r, column=15).value
        floor_val = ws02_data.cell(row=r, column=16).value
        final_score = ws02_data.cell(row=r, column=17).value
        risk_tier = ws02_data.cell(row=r, column=18).value
        action = ws02_data.cell(row=r, column=19).value
        
        def _safe_float(val):
            if val is None or val == "-" or val == "":
                return None
            try:
                return float(val)
            except (ValueError, TypeError):
                return None

        customers.append({
            "row": r,
            "row_num": row_num,
            "id": cid,
            "name": name,
            "nid": nid,
            "id_status": id_status,
            "cbi_color": cbi_color,
            "chq_cnt": int(chq_cnt),
            "fund_amt": int(fund_amt),
            "fund_share": fund_share,
            "in_transit": int(in_transit),
            "bounced": int(bounced),
            "cleared_amt": int(cleared_amt),
            "pers_cat": pers_cat,
            "pers_periods": int(pers_periods),
            "raw_score": _safe_float(raw_score),
            "applied_floor": floor_val,
            "risk_score": _safe_float(final_score),
            "risk_tier": risk_tier,
            "action": action
        })
    
    print(f"\n[AC 1 - Deduplication Integrity]:")
    print(f"  Total canonical customer rows in Sheet 02: {len(customers)} (Expected 48)")
    assert len(customers) == 48, f"Expected 48 customers, found {len(customers)}"
    
    # Check deduplicated sums
    tot_fund = sum(c["fund_amt"] for c in customers)
    tot_cheques = sum(c["chq_cnt"] for c in customers)
    tot_in_transit = sum(c["in_transit"] for c in customers)
    tot_bounced = sum(c["bounced"] for c in customers)
    tot_cleared = sum(c["cleared_amt"] for c in customers)
    
    print(f"  Total fund cheques amount: {tot_fund:,} Rials (Expected: 483,325,000,000)")
    print(f"  Total fund cheques count: {tot_cheques} (Expected: 147)")
    print(f"  Total in-transit bank amount: {tot_in_transit:,} Rials (Expected: 4,466,069,469,454)")
    print(f"  Total bounced bank amount: {tot_bounced:,} Rials (Expected: 231,951,000,000)")
    print(f"  Total cleared bank amount: {tot_cleared:,} Rials (Expected: 1,024,305,185,797)")
    
    assert tot_fund == 483_325_000_000, "Fund amount mismatch!"
    assert tot_cheques == 147, "Cheques count mismatch!"
    assert tot_in_transit == 4_466_069_469_454, "In-transit mismatch!"
    assert tot_bounced == 231_951_000_000, "Bounced mismatch!"
    assert tot_cleared == 1_024_305_185_797, "Cleared mismatch!"
    print("  => AC 1 & AC 2 Totals: PASS")
    
    # Check Sheet 02 row 52 formulas
    print(f"\n[AC 1 & AC 2 - Sheet 02 Summary Row 52 Formulas]:")
    print(f"  G52 (cheque count): {ws02.cell(row=52, column=7).value}")
    print(f"  H52 (fund amount): {ws02.cell(row=52, column=8).value}")
    print(f"  J52 (in-transit): {ws02.cell(row=52, column=10).value}")
    print(f"  K52 (bounced): {ws02.cell(row=52, column=11).value}")
    print(f"  L52 (cleared): {ws02.cell(row=52, column=12).value}")
    assert ws02.cell(row=52, column=8).value == "=SUM(H4:H51)"
    assert ws02.cell(row=52, column=10).value == "=SUM(J4:J51)"
    assert ws02.cell(row=52, column=11).value == "=SUM(K4:K51)"
    assert ws02.cell(row=52, column=12).value == "=SUM(L4:L51)"
    print("  => Sheet 02 Summary Formulas: PASS")
    
    # 4. Sheet 07 Detail Cheques Audit
    ws07 = wb["07_چکهای_تفصیلی"]
    ws07_data = wb_data["07_چکهای_تفصیلی"]
    cheques = []
    for r in range(2, 149):
        c_amt = ws07_data.cell(row=r, column=12).value
        if c_amt is not None:
            cheques.append(int(c_amt))
    print(f"\n[AC 2 - Sheet 07 Detail Cheques]:")
    print(f"  Cheque count: {len(cheques)} (Expected 147)")
    print(f"  Cheque sum: {sum(cheques):,} (Expected 483,325,000,000)")
    assert len(cheques) == 147
    assert sum(cheques) == 483_325_000_000
    formula_07 = ws07.cell(row=149, column=12).value
    print(f"  Row 149 formula: {formula_07}")
    assert formula_07 == "=SUM(L2:L148)"
    print("  => AC 2 Sheet 07: PASS")
    
    # 5. AC 3 Identity Precision (0933387075 & Zahra Bahrami Pouya)
    heshmati = [c for c in customers if "حسین حشمتی" in str(c["name"])]
    alipour = [c for c in customers if "امیرحسین علیپور" in str(c["name"])]
    h_fund = heshmati[0]['fund_amt'] if heshmati else 0
    a_fund = alipour[0]['fund_amt'] if alipour else 0
    h_nid = heshmati[0]['nid'] if heshmati else None
    a_nid = alipour[0]['nid'] if alipour else None
    print(f"\n[AC 3 - Disambiguation 0933387075]:")
    print(f"  Hossein Heshmati: count={len(heshmati)}, NID={h_nid}, Fund={h_fund:,}")
    print(f"  Amirhossein Alipour: count={len(alipour)}, NID={a_nid}, Fund={a_fund:,}")
    assert len(heshmati) == 1
    assert len(alipour) == 1
    assert heshmati[0]["nid"] == "0933387075"
    assert alipour[0]["nid"] == "0933387075"
    assert heshmati[0]["id"] != alipour[0]["id"]
    print("  => Disambiguation: PASS")
    
    # Check Sheet 03 Zahra Bahrami Pouya
    ws03_data = wb_data["03_رخدادهای_امروز"]
    zahra_events = []
    for r in range(4, 30):
        name = ws03_data.cell(row=r, column=2).value
        nid = ws03_data.cell(row=r, column=3).value
        ev_type = ws03_data.cell(row=r, column=4).value
        d_inf = ws03_data.cell(row=r, column=7).value
        d_bnc = ws03_data.cell(row=r, column=10).value
        conf = ws03_data.cell(row=r, column=12).value
        if name and "زهرا بهرامی" in str(name):
            zahra_events.append({
                "row": r, "name": name, "nid": nid, "event": ev_type,
                "delta_in_flight": d_inf, "delta_bounced": d_bnc, "confidence": conf
            })
    print(f"\n[AC 3 - Zahra Bahrami Pouya Events]: count={len(zahra_events)}")
    for ze in zahra_events:
        print(f"  Row {ze['row']}: {ze['name']} ({ze['nid']}) - {ze['event']} | d_inf={ze['delta_in_flight']} | d_bnc={ze['delta_bounced']} | conf={ze['confidence']}")
    assert len(zahra_events) >= 1
    assert any(abs(int(ze["delta_bounced"])) == 3_500_000_000 for ze in zahra_events)
    print("  => AC 3 Zahra Bahrami Pouya Transition: PASS")
    
    # 6. AC 4 Risk Floors Verification
    print(f"\n[AC 4 - Mandatory Risk Floors Check on all 48 customers]:")
    floor_violations = []
    for c in customers:
        bounced = c["bounced"]
        periods = c["pers_periods"]
        is_persistent = periods >= 3
        is_new = 1 <= periods <= 2
        score = c["risk_score"]
        name = c["name"]
        
        if score is None:
            if bounced > 0:
                floor_violations.append((name, bounced, periods, None, "numeric score"))
            continue
            
        if bounced > 50_000_000_000 and is_persistent and score < 85:
            floor_violations.append((name, bounced, periods, score, 85))
        elif bounced > 20_000_000_000 and is_persistent and score < 78:
            floor_violations.append((name, bounced, periods, score, 78))
        elif bounced > 10_000_000_000 and is_persistent and score < 72:
            floor_violations.append((name, bounced, periods, score, 72))
        elif bounced > 5_000_000_000 and is_persistent and score < 65:
            floor_violations.append((name, bounced, periods, score, 65))
        elif bounced > 0 and is_persistent and score < 45:
            floor_violations.append((name, bounced, periods, score, 45))
        elif bounced > 0 and is_new and score < 35:
            floor_violations.append((name, bounced, periods, score, 35))
            
    print(f"  Total floor violations: {len(floor_violations)}")
    if floor_violations:
        for v in floor_violations:
            print(f"    VIOLATION: {v[0]} bounced={v[1]:,} periods={v[2]} score={v[3]} (required >={v[4]})")
    assert len(floor_violations) == 0, f"Found {len(floor_violations)} floor violations!"
    print("  => AC 4 Risk Floors: PASS")
    
    # 7. Sheet 09 Dual-Audit Invariants
    ws09 = wb["09_ممیزی"]
    print(f"\n[AC 4 - Sheet 09 Dual-Audit Formulas (14 automated checks)]:")
    for r in range(8, 22):
        row_idx = ws09.cell(row=r, column=1).value
        code = ws09.cell(row=r, column=2).value
        title = ws09.cell(row=r, column=3).value
        target = ws09.cell(row=r, column=5).value
        exp = ws09.cell(row=r, column=6).value
        formula = ws09.cell(row=r, column=8).value
        print(f"  {code}: {title}")
        print(f"    Target: {target}")
        print(f"    Expected: {exp}")
        print(f"    Formula: {formula}")
        assert code == f"AUD-{r-7:02d}", f"Expected AUD-{r-7:02d}, got {code}"
        assert formula.startswith("="), f"Formula missing =: {formula}"
        assert '"PASS"' in formula and '"FAIL"' in formula, f"Formula doesn't return PASS/FAIL: {formula}"
    print("  => AC 4 Dual-Audit Formulas: PASS")
    
    print(f"\n=======================================================")
    print(f"ALL AUDIT CRITERIA SATISFIED FOR {os.path.basename(filepath)}")
    print(f"=======================================================\n")

if __name__ == "__main__":
    audit_file(ROOT_FILE)
    audit_file(DESKTOP_FILE)
